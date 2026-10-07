import json
from pathlib import Path

import pytest

from centelha_api.models import Capitulo, Edicao, EstadoCapitulo, Obra, Segmento, TipoSegmento
from centelha_api.pipeline.ingestao import leitores
from centelha_api.pipeline.ingestao.cli import main as cli
from centelha_api.pipeline.ingestao.estrutura import estruturar, resumo
from centelha_api.pipeline.ingestao.importar import ImportacaoRecusada, importar

# Texto sintético no formato de O Livro dos Espíritos (não é citação da obra).
LE_EXEMPLO = """\
INTRODUÇÃO

Texto de abertura da introdução, com um parágrafo cor-
rido que quebra a linha no meio.

LIVRO PRIMEIRO

CAPÍTULO I

DO PRIMEIRO TEMA

1. Primeira pergunta de exemplo?

“Primeira resposta de exemplo.”

2. Segunda pergunta de exemplo?

“Segunda resposta, primeiro parágrafo.

“Segunda resposta, segundo parágrafo.”

Comentário do autor sobre a segunda questão.

a) Subquestão ligada à segunda?

“Resposta da subquestão.”

CAPÍTULO II — DO SEGUNDO TEMA

4. Quarta pergunta, a terceira falta de propósito?

“Quarta resposta.”
"""


@pytest.fixture
def paragrafos_le(tmp_path):
    arquivo = tmp_path / "le.txt"
    arquivo.write_text(LE_EXEMPLO, encoding="utf-8")
    return leitores.ler(arquivo)


def test_leitor_txt_junta_linhas_e_hifenizacao(paragrafos_le):
    assert paragrafos_le[1] == (
        "Texto de abertura da introdução, com um parágrafo corrido que quebra a linha no meio."
    )


def test_estrutura_perfil_perguntas(paragrafos_le):
    caps = estruturar(paragrafos_le, "perguntas")
    assert [c.titulo for c in caps] == [
        "INTRODUÇÃO",
        "LIVRO PRIMEIRO — CAPÍTULO I — DO PRIMEIRO TEMA",
        "LIVRO PRIMEIRO — CAPÍTULO II — DO SEGUNDO TEMA",
    ]
    tipos = [(s.tipo.value, s.numero_questao, s.subquestao) for s in caps[1].segmentos]
    assert tipos == [
        ("titulo", None, None),
        ("pergunta", 1, None),
        ("resposta", 1, None),
        ("pergunta", 2, None),
        ("resposta", 2, None),
        ("resposta", 2, None),
        ("comentario", 2, None),
        ("pergunta", 2, "a"),
        ("resposta", 2, "a"),
    ]
    assert caps[1].segmentos[1].texto == "Primeira pergunta de exemplo?"


def test_resumo_aponta_questao_faltando(paragrafos_le):
    r = resumo(estruturar(paragrafos_le, "perguntas"))
    assert r["capitulos"] == 3
    assert (r["primeira_questao"], r["ultima_questao"]) == (1, 4)
    assert r["questoes_faltando"] == [3]
    assert r["caracteres"] > 0


def test_perfil_generico_nao_inventa_perguntas(paragrafos_le):
    caps = estruturar(paragrafos_le, "generico")
    tipos = {s.tipo for c in caps for s in c.segmentos}
    assert tipos == {TipoSegmento.TITULO, TipoSegmento.PARAGRAFO}


def test_nota_vira_segmento_nota_sem_quebrar_a_resposta():
    paragrafos = [
        "CAPÍTULO I",
        "1. Pergunta?",
        "“Resposta.”",
        leitores.Nota("Nota de Kardec."),
        "“Continuação da resposta.”",
    ]
    caps = estruturar(paragrafos, "perguntas")
    tipos = [(s.tipo.value, s.numero_questao) for s in caps[0].segmentos[1:]]
    assert tipos == [("pergunta", 1), ("resposta", 1), ("nota", 1), ("resposta", 1)]


def test_titulos_do_pdf_da_feb():
    paragrafos = [
        "I N T R O D U Ç Ã O",
        "ao estudo da",
        "DOUTRINA ESPÍRITA",
        "Texto da introdução.",
        "capítulo xxviii",
        "Coletânea de preces espíritas",
        "Preâmbulo",
        "1. Os Espíritos hão dito sempre.",
    ]
    caps = estruturar(paragrafos, "generico")
    # Versalete normalizado; "Preâmbulo" depois do primeiro capítulo é seção, não capítulo.
    assert [c.titulo for c in caps] == [
        "INTRODUÇÃO ao estudo da DOUTRINA ESPÍRITA",
        "Capítulo XXVIII — Coletânea de preces espíritas",
    ]
    assert [s.texto for s in caps[1].segmentos[1:]] == [
        "Preâmbulo",
        "1. Os Espíritos hão dito sempre.",
    ]


def test_texto_antes_do_primeiro_capitulo_vira_abertura():
    caps = estruturar(["Folha de rosto.", "Capítulo 1", "Texto."], "generico")
    assert [c.titulo for c in caps] == ["Abertura", "Capítulo 1"]


def _paragrafos_exemplo():
    return ["CAPÍTULO I", "O TÍTULO", "1. Uma pergunta?", "“Uma resposta.”"]


def test_leitor_docx(tmp_path):
    import docx

    doc = docx.Document()
    for p in _paragrafos_exemplo():
        doc.add_paragraph(p)
    doc.add_paragraph("   ")
    caminho = tmp_path / "x.docx"
    doc.save(caminho)
    assert leitores.ler(caminho) == _paragrafos_exemplo()


def test_leitor_epub(tmp_path):
    from ebooklib import epub

    livro = epub.EpubBook()
    livro.set_identifier("teste")
    livro.set_title("Teste")
    livro.set_language("pt-BR")
    cap = epub.EpubHtml(title="c1", file_name="c1.xhtml", lang="pt-BR")
    cap.content = (
        "<html><body><h1>CAPÍTULO I</h1><h2>O TÍTULO</h2>"
        "<p>1. Uma pergunta?</p><blockquote><p>“Uma resposta.”</p></blockquote>"
        "</body></html>"
    )
    livro.add_item(cap)
    livro.add_item(epub.EpubNcx())
    livro.add_item(epub.EpubNav())
    livro.spine = [cap]
    caminho = tmp_path / "x.epub"
    epub.write_epub(str(caminho), livro)
    assert leitores.ler(caminho) == _paragrafos_exemplo()


def test_leitor_pdf_ignora_numero_de_pagina(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    pagina = doc.new_page()
    y = 72
    for p in _paragrafos_exemplo():
        pagina.insert_text((72, y), p, fontname="helv")
        y += 40
    pagina.insert_text((290, 820), "12", fontname="helv")  # no pé da página
    caminho = tmp_path / "x.pdf"
    doc.save(caminho)
    # Fonte base do PDF não tem aspas curvas; o leitor só não pode perder o texto.
    lidos = leitores.ler(caminho)
    assert lidos[:3] == _paragrafos_exemplo()[:3]
    assert "12" not in lidos
    assert len(lidos) == 4


def _pdf_como_o_da_feb(caminho):
    """Três páginas com o que os PDFs da FEB (#2) trazem em volta do texto.

    Cabeçalho corrido e número de página; parágrafo que vira a página; chamada de nota
    em sobrescrito; nota da editora (N.E.) que continua no rodapé da página seguinte;
    nota de Kardec no estilo do LE, depois de "________", com chamada "(1)"; e "(85)",
    remissão à questão 85, que é texto.
    """
    import pymupdf

    doc = pymupdf.open()
    for _ in range(3):
        doc.new_page()
    # Página criada antes de outra new_page() perde a referência ao documento no PyMuPDF.
    paginas = [doc[n] for n in range(3)]
    for n, pagina in enumerate(paginas, start=1):
        pagina.insert_text((250, 40), "O Livro de Teste", fontsize=10)
        pagina.insert_text((290, 820), str(n), fontsize=10)

    p1, p2, p3 = paginas
    p1.insert_text((72, 120), "Primeiro parágrafo.", fontsize=12)
    p1.insert_text((72, 160), "Este parágrafo continua na página", fontsize=12)
    p1.insert_text((262, 156), "1", fontsize=6)
    p1.insert_text((72, 700), "1 N.E. de 1947: texto da editora, que tem direito", fontsize=8)
    p1.insert_text((72, 710), "próprio e não pode entrar", fontsize=8)

    p2.insert_text((72, 120), "seguinte e termina aqui.", fontsize=12)
    p2.insert_text((72, 700), "na narração nem como nota.", fontsize=8)

    p3.insert_text((72, 120), "Texto de Kardec, ver a questão (85). Fim (1).", fontsize=12)
    p3.insert_text((72, 600), "________", fontsize=12)
    p3.insert_text((72, 620), "(1) Nota de Kardec.", fontsize=11)
    doc.save(caminho)
    return caminho


def test_pdf_tira_o_que_nao_e_texto_da_obra(tmp_path):
    lidos = leitores.ler(_pdf_como_o_da_feb(tmp_path / "feb.pdf"))
    assert lidos == [
        "Primeiro parágrafo.",
        "Este parágrafo continua na página seguinte e termina aqui.",
        "Texto de Kardec, ver a questão (85). Fim.",
        "Nota de Kardec.",
    ]
    assert [isinstance(p, leitores.Nota) for p in lidos] == [False, False, False, True]


def test_pdf_nunca_deixa_passar_nota_da_editora(tmp_path):
    """Direitos (#2): a nota da FEB não entra nem grudada numa nota de Kardec.

    Se a leitura juntar as duas (aconteceu: o número da nota sumiu junto com as chamadas
    em sobrescrito), a nota inteira sai. Perder a de Kardec volta na revisão; vazar texto
    da editora é publicar o que não tem direito.
    """
    import pymupdf

    doc = pymupdf.open()
    pagina = doc.new_page()
    # O corpo do texto é o tamanho de letra com mais caracteres na página.
    pagina.insert_text(
        (72, 120), "Texto corrido, mais longo que as notas para mandar no", fontsize=12
    )
    pagina.insert_text((72, 135), "corpo do texto, como numa página de verdade.", fontsize=12)
    pagina.insert_text((72, 700), "Nota de Allan Kardec: nota legítima.", fontsize=8)
    pagina.insert_text((72, 720), "N.E.: Ver Nota Explicativa, p. 371.", fontsize=8)
    caminho = tmp_path / "grudada.pdf"
    doc.save(caminho)
    lidos = leitores.ler(caminho)
    assert lidos == [
        "Texto corrido, mais longo que as notas para mandar no corpo do texto, como numa "
        "página de verdade."
    ]


def test_pdf_faixa_de_paginas(tmp_path):
    caminho = _pdf_como_o_da_feb(tmp_path / "feb.pdf")
    assert leitores.ler(caminho, (2, 3))[0] == "seguinte e termina aqui."
    with pytest.raises(ValueError, match="fora"):
        leitores.ler(caminho, (2, 9))
    txt = tmp_path / "x.txt"
    txt.write_text("Texto.", encoding="utf-8")
    with pytest.raises(ValueError, match="só vale para PDF"):
        leitores.ler(txt, (1, 1))


def test_pdf_titulos(tmp_path):
    """Título em duas linhas vira um; subtítulo no mesmo bloco do texto se separa; título
    em versalete depois de parágrafo aberto não é continuação dele."""
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    doc.new_page()
    p1, p2 = doc[0], doc[1]
    p1.insert_text((100, 200), "Do mundo espírita ou mundo dos", fontsize=18)
    p1.insert_text((200, 222), "Espíritos", fontsize=18)
    p1.insert_text((150, 300), "Uma realeza terrestre", fontsize=14)
    p1.insert_text((72, 318), "8. Quem melhor do que eu pode compreender a verdade", fontsize=12)
    p1.insert_text((72, 333), "desta palavra, que fica aberta", fontsize=12)
    p2.insert_text((250, 150), "capítulo ii", fontsize=11)
    p2.insert_text(
        (72, 200), "Texto do capítulo, longo o bastante para mandar no corpo.", fontsize=12
    )
    caminho = tmp_path / "titulos.pdf"
    doc.save(caminho)
    assert leitores.ler(caminho) == [
        "Do mundo espírita ou mundo dos Espíritos",
        "Uma realeza terrestre",
        "8. Quem melhor do que eu pode compreender a verdade desta palavra, que fica aberta",
        "capítulo ii",
        "Texto do capítulo, longo o bastante para mandar no corpo.",
    ]


def test_limpar_hifen_de_pronome_e_caracteres_do_pdf():
    # Pronome depois de vogal acentuada mantém o hífen; o resto é hifenização de linha.
    assert leitores.limpar("avaliá-\nla e fazê-\nlo, fra-\nse") == "avaliá-la e fazê-lo, frase"
    # \x03 é o espaço antes da referência bíblica; U+00AD, o hífen discricionário.
    assert leitores.limpar("ponto.\x03(Mateus)") == "ponto. (Mateus)"
    assert leitores.limpar("pre\u00ad\ncisamos \u00adisso") == "precisamos isso"


def test_cortar_em():
    paragrafos = ["Conclusão.", "Santo Agostinho.", "Nota Especial n°1, da editora."]
    assert leitores.cortar_em(paragrafos, "Nota Especial") == paragrafos[:2]
    with pytest.raises(ValueError, match="nenhum"):
        leitores.cortar_em(paragrafos, "Nota Espacial")


def test_formato_nao_suportado(tmp_path):
    with pytest.raises(ValueError, match="formato"):
        leitores.ler(tmp_path / "x.odt")


def _edicao(session):
    obra = Obra(
        slug="o-livro-dos-espiritos",
        autor="Allan Kardec",
        titulo_original="Le Livre des Esprits",
        idioma_original="fr",
        sigla="LE",
    )
    edicao = Edicao(obra=obra, idioma="pt-BR", titulo="O Livro dos Espíritos", fonte="teste")
    session.add(edicao)
    session.commit()
    return edicao


def test_importar_grava_capitulos_e_segmentos(session, paragrafos_le):
    edicao = _edicao(session)
    importar(session, edicao, estruturar(paragrafos_le, "perguntas"))
    session.commit()
    caps = session.query(Capitulo).order_by(Capitulo.ordem).all()
    assert [c.referencia_canonica for c in caps] == ["LE-C001", "LE-C002", "LE-C003"]
    assert {c.estado for c in caps} == {EstadoCapitulo.IMPORTADO}
    sub = session.query(Segmento).filter_by(subquestao="a", tipo=TipoSegmento.PERGUNTA).one()
    assert sub.numero_questao == 2


def test_reimportar_exige_substituir_e_protege_revisados(session, paragrafos_le):
    edicao = _edicao(session)
    caps = estruturar(paragrafos_le, "perguntas")
    importar(session, edicao, caps)
    session.commit()
    with pytest.raises(ImportacaoRecusada, match="substituir"):
        importar(session, edicao, caps)
    importar(session, edicao, caps[:1], substituir=True)
    session.commit()
    assert session.query(Capitulo).count() == 1

    session.query(Capitulo).one().estado = EstadoCapitulo.TEXTO_REVISADO
    session.commit()
    with pytest.raises(ImportacaoRecusada, match="revisados"):
        importar(session, edicao, caps, substituir=True)


def test_cli_resumo_e_json(tmp_path, capsys):
    fonte = tmp_path / "le.txt"
    fonte.write_text(LE_EXEMPLO, encoding="utf-8")
    saida = tmp_path / "estrutura.json"
    assert cli([str(fonte), "--perfil", "perguntas", "--json", str(saida)]) == 0
    assert json.loads(capsys.readouterr().out)["questoes_faltando"] == [3]
    assert json.loads(Path(saida).read_text(encoding="utf-8"))[1]["segmentos"][1]["tipo"] == (
        "pergunta"
    )


def test_cli_paginas_e_cortar_em(tmp_path, capsys):
    fonte = _pdf_como_o_da_feb(tmp_path / "feb.pdf")
    saida = tmp_path / "estrutura.json"
    args = [str(fonte), "--paginas", "1-3", "--cortar-em", "Texto de Kardec", "--json", str(saida)]
    assert cli(args) == 0
    textos = [
        s["texto"] for c in json.loads(saida.read_text(encoding="utf-8")) for s in c["segmentos"]
    ]
    assert textos[-1] == "Este parágrafo continua na página seguinte e termina aqui."
    with pytest.raises(SystemExit):
        cli([str(fonte), "--paginas", "três"])


# Texto sintético no formato do original francês (não é citação da obra).
LE_EXEMPLO_FR = """\
INTRODUCTION À L’ÉTUDE DE LA DOCTRINE SPIRITE

Texte d’ouverture de l’introduction.

LIVRE PREMIER

CHAPITRE PREMIER

DU PREMIER SUJET

1. Première question d’exemple ?

« Première réponse d’exemple. »

Commentaire de l’auteur.

CHAPITRE II. — DU SECOND SUJET

2. Deuxième question ?

« Deuxième réponse. »

PREMIÈRE PARTIE — DOCTRINE

CHAPITRE III

Un paragraphe narré.
"""


def test_estrutura_reconhece_titulos_em_frances(tmp_path):
    arquivo = tmp_path / "le-fr.txt"
    arquivo.write_text(LE_EXEMPLO_FR, encoding="utf-8")
    caps = estruturar(leitores.ler(arquivo), "perguntas")
    assert [c.titulo for c in caps] == [
        "INTRODUCTION À L’ÉTUDE DE LA DOCTRINE SPIRITE",
        "LIVRE PREMIER — CHAPITRE PREMIER — DU PREMIER SUJET",
        "LIVRE PREMIER — CHAPITRE II — DU SECOND SUJET",
        "PREMIÈRE PARTIE — DOCTRINE — CHAPITRE III",
    ]
    tipos = [(s.tipo.value, s.numero_questao) for s in caps[1].segmentos]
    assert tipos == [
        ("titulo", None),
        ("pergunta", 1),
        ("resposta", 1),
        ("comentario", 1),
    ]
